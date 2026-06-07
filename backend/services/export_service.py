import io
from typing import Optional
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, ListFlowable, ListItem
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import json


def generate_pdf(meeting_data: dict, title: Optional[str] = "Meeting Summary") -> bytes:
    """Generate a professional PDF from meeting notes."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
    )

    styles = getSampleStyleSheet()
    accent = colors.HexColor("#2563EB")

    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Title"],
        textColor=accent,
        fontSize=22,
        spaceAfter=6,
    )
    heading_style = ParagraphStyle(
        "CustomHeading",
        parent=styles["Heading2"],
        textColor=accent,
        fontSize=13,
        spaceBefore=14,
        spaceAfter=4,
    )
    body_style = ParagraphStyle(
        "CustomBody",
        parent=styles["Normal"],
        fontSize=10,
        leading=16,
        spaceAfter=4,
    )

    story = []

    story.append(Paragraph(title, title_style))
    story.append(HRFlowable(width="100%", thickness=1, color=accent))
    story.append(Spacer(1, 12))

    # Summary
    story.append(Paragraph("Meeting Summary", heading_style))
    story.append(Paragraph(meeting_data.get("summary", "No summary available."), body_style))
    story.append(Spacer(1, 8))

    # Action Items
    action_items = meeting_data.get("action_items", [])
    if action_items:
        story.append(Paragraph("Action Items", heading_style))
        items = []
        for item in action_items:
            task = item.get("task", "")
            assignee = item.get("assignee") or "Unassigned"
            deadline = item.get("deadline") or "No deadline"
            text = f"<b>{task}</b> — {assignee} ({deadline})"
            items.append(ListItem(Paragraph(text, body_style), bulletColor=accent))
        story.append(ListFlowable(items, bulletType="bullet"))
        story.append(Spacer(1, 8))

    # Decisions
    decisions = meeting_data.get("decisions", [])
    if decisions:
        story.append(Paragraph("Key Decisions", heading_style))
        items = [ListItem(Paragraph(d, body_style), bulletColor=accent) for d in decisions]
        story.append(ListFlowable(items, bulletType="bullet"))
        story.append(Spacer(1, 8))

    # Key Points
    key_points = meeting_data.get("key_points", [])
    if key_points:
        story.append(Paragraph("Key Discussion Points", heading_style))
        items = [ListItem(Paragraph(kp, body_style), bulletColor=accent) for kp in key_points]
        story.append(ListFlowable(items, bulletType="bullet"))

    doc.build(story)
    return buffer.getvalue()


def generate_docx(meeting_data: dict, title: Optional[str] = "Meeting Summary") -> bytes:
    """Generate an editable DOCX from meeting notes."""
    doc = Document()

    # Title
    heading = doc.add_heading(title, 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    doc.add_paragraph()

    # Summary
    doc.add_heading("Meeting Summary", level=1)
    doc.add_paragraph(meeting_data.get("summary", "No summary available."))

    # Action Items
    action_items = meeting_data.get("action_items", [])
    if action_items:
        doc.add_heading("Action Items", level=1)
        for item in action_items:
            task = item.get("task", "")
            assignee = item.get("assignee") or "Unassigned"
            deadline = item.get("deadline") or "No deadline"
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(task)
            run.bold = True
            p.add_run(f" — {assignee} ({deadline})")

    # Decisions
    decisions = meeting_data.get("decisions", [])
    if decisions:
        doc.add_heading("Key Decisions", level=1)
        for d in decisions:
            doc.add_paragraph(d, style="List Bullet")

    # Key Points
    key_points = meeting_data.get("key_points", [])
    if key_points:
        doc.add_heading("Key Discussion Points", level=1)
        for kp in key_points:
            doc.add_paragraph(kp, style="List Bullet")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def generate_txt(meeting_data: dict, title: Optional[str] = "Meeting Summary") -> bytes:
    """Generate a plain text version of meeting notes."""
    lines = []
    sep = "=" * 60

    lines.append(sep)
    lines.append(f"  {title.upper()}")
    lines.append(sep)
    lines.append("")

    lines.append("SUMMARY")
    lines.append("-" * 40)
    lines.append(meeting_data.get("summary", "No summary available."))
    lines.append("")

    action_items = meeting_data.get("action_items", [])
    if action_items:
        lines.append("ACTION ITEMS")
        lines.append("-" * 40)
        for item in action_items:
            task = item.get("task", "")
            assignee = item.get("assignee") or "Unassigned"
            deadline = item.get("deadline") or "No deadline"
            lines.append(f"  • {task} — {assignee} ({deadline})")
        lines.append("")

    decisions = meeting_data.get("decisions", [])
    if decisions:
        lines.append("KEY DECISIONS")
        lines.append("-" * 40)
        for d in decisions:
            lines.append(f"  • {d}")
        lines.append("")

    key_points = meeting_data.get("key_points", [])
    if key_points:
        lines.append("KEY DISCUSSION POINTS")
        lines.append("-" * 40)
        for kp in key_points:
            lines.append(f"  • {kp}")
        lines.append("")

    lines.append(sep)

    return "\n".join(lines).encode("utf-8")
